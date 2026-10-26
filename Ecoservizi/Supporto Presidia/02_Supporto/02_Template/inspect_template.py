from docx import Document
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = Document(r'temp_extract/21_Stress lavoro correlato/21_Stress lavoro correlato_UFFICI.docx')
print('=== TEMPLATE: UFFICI ===')
print(f'TOTAL TABLES: {len(doc.tables)}')

for table_idx, table in enumerate(doc.tables):
    rows_count = len(table.rows)
    cols_count = len(table.columns)
    separator = "="*80
    print(f'\n{separator}')
    print(f'TABLE {table_idx}: {rows_count} rows x {cols_count} columns')
    print(f'{separator}')
    
    for row_idx in range(min(4, rows_count)):
        print(f'\nROW {row_idx}:')
        for col_idx, cell in enumerate(table.rows[row_idx].cells):
            text = cell.text.strip()[:150]
            print(f'  [{col_idx}]: {text}')
